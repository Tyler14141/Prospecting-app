"""
Append Competitor Summary and Competitor Penetration tabs to the workbook.
Also writes competitors_data.json for the dashboard.
"""

import json
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

from competitors import COMPETITORS, STATE_PENETRATION, vendor_state_total
from competitor_customers import CUSTOMERS, by_vendor as customers_by_vendor
from data import STATE_COUNTS, STATE_NAMES, STATES

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(HERE, "..", "TAM"))
XLSX_PATH = os.path.join(OUT_DIR, "Local_Government_TAM.xlsx")
JSON_PATH = os.path.join(OUT_DIR, "competitors_data.json")

ARIAL = "Arial"
FILL_HEADER  = PatternFill("solid", fgColor="0F172A")
FILL_TOTAL   = PatternFill("solid", fgColor="E5E7EB")
FONT_HEADER  = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
FONT_BODY    = Font(name=ARIAL, size=10)
FONT_INPUT   = Font(name=ARIAL, size=10, color="1F4E79")
FONT_FORMULA = Font(name=ARIAL, size=10, color="000000")
FONT_TOTAL   = Font(name=ARIAL, size=10, bold=True)
FONT_TITLE   = Font(name=ARIAL, size=14, bold=True)

THIN = Side(style="thin", color="CBD5E1")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUM_INT = '#,##0;(#,##0);"-"'
NUM_PCT = '0.00%;(0.00%);"-"'


def build_summary_tab(wb):
    ws = wb.create_sheet("Competitor Summary")
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="Competitor Summary").font = FONT_TITLE
    ws.cell(row=2, column=1, value="Per-vendor profile across the U.S. local-government software market.").font = FONT_BODY

    headers = ["Vendor", "HQ", "Customers (published)", "Customers (allocated)",
               "Calibration Δ", "Focus", "Size Target", "Top State", "#2 State",
               "Notes"]
    for i, h in enumerate(headers):
        c = ws.cell(row=4, column=i + 1, value=h)
        c.font = FONT_HEADER; c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER_ALL
    ws.row_dimensions[4].height = 30

    row = 5
    for vendor, meta in COMPETITORS.items():
        sp = STATE_PENETRATION[vendor]
        sorted_states = sorted(sp.items(), key=lambda x: -x[1])
        top_st = f"{sorted_states[0][0]} ({sorted_states[0][1]})"
        snd_st = f"{sorted_states[1][0]} ({sorted_states[1][1]})"
        allocated = vendor_state_total(vendor)
        dev = (allocated - meta["total"]) / meta["total"]

        cells = [
            (1, vendor, FONT_TOTAL),
            (2, f"{meta['hq']} ({meta['hq_state_full']})", FONT_BODY),
            (3, meta["total"], FONT_INPUT),
            (4, allocated, FONT_FORMULA),
            (5, dev, FONT_FORMULA),
            (6, meta["focus"], FONT_BODY),
            (7, meta["size_target"], FONT_BODY),
            (8, top_st, FONT_BODY),
            (9, snd_st, FONT_BODY),
            (10, meta["notes"], FONT_BODY),
        ]
        for col, val, font in cells:
            c = ws.cell(row=row, column=col, value=val)
            c.font = font
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER_ALL
        ws.cell(row=row, column=3).number_format = NUM_INT
        ws.cell(row=row, column=4).number_format = NUM_INT
        ws.cell(row=row, column=5).number_format = NUM_PCT
        ws.row_dimensions[row].height = 60
        row += 1

    widths = [22, 22, 14, 14, 11, 32, 22, 14, 14, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B5"


def build_penetration_tab(wb):
    ws = wb.create_sheet("Competitor Penetration")
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="Competitor Penetration — State × Vendor matrix").font = FONT_TITLE
    ws.cell(row=2, column=1, value=("Cells = customer count by state. Combined Penetration % = "
                                     "total tracked-vendor customers ÷ (counties + munis + townships) "
                                     "in that state.")).font = FONT_BODY

    vendors = list(COMPETITORS.keys())
    headers = ["State", "Counties", "Munis", "Townships", "Addressable (C+M+T)"] + vendors + ["Combined", "Combined Pen %"]
    for i, h in enumerate(headers):
        c = ws.cell(row=4, column=i + 1, value=h)
        c.font = FONT_HEADER; c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER_ALL
    ws.row_dimensions[4].height = 36

    start_row = 5
    for i, st in enumerate(STATES):
        row = start_row + i
        c, m, t, _ = STATE_COUNTS[st]
        addressable = c + m + t

        ws.cell(row=row, column=1, value=st).font = FONT_TOTAL
        for col, val in [(2, c), (3, m), (4, t)]:
            cc = ws.cell(row=row, column=col, value=val)
            cc.font = FONT_INPUT; cc.number_format = NUM_INT
        cc = ws.cell(row=row, column=5,
                     value=f"=B{row}+C{row}+D{row}")
        cc.font = FONT_FORMULA; cc.number_format = NUM_INT

        for j, v in enumerate(vendors):
            cc = ws.cell(row=row, column=6 + j, value=STATE_PENETRATION[v][st])
            cc.font = FONT_INPUT; cc.number_format = NUM_INT

        # Combined
        first = get_column_letter(6)
        last = get_column_letter(6 + len(vendors) - 1)
        cc = ws.cell(row=row, column=6 + len(vendors),
                     value=f"=SUM({first}{row}:{last}{row})")
        cc.font = FONT_FORMULA; cc.number_format = NUM_INT
        # Pen %
        cc = ws.cell(row=row, column=7 + len(vendors),
                     value=f"={get_column_letter(6 + len(vendors))}{row}/E{row}")
        cc.font = FONT_FORMULA; cc.number_format = NUM_PCT

    end_row = start_row + len(STATES) - 1

    # Totals
    total_row = end_row + 1
    ws.cell(row=total_row, column=1, value="US TOTAL").font = FONT_TOTAL
    ws.cell(row=total_row, column=1).fill = FILL_TOTAL
    last_col = 7 + len(vendors)
    for col in range(2, last_col):
        col_letter = get_column_letter(col)
        cc = ws.cell(row=total_row, column=col,
                     value=f"=SUM({col_letter}{start_row}:{col_letter}{end_row})")
        cc.font = FONT_TOTAL; cc.fill = FILL_TOTAL; cc.number_format = NUM_INT
    cc = ws.cell(row=total_row, column=last_col,
                 value=f"={get_column_letter(6 + len(vendors))}{total_row}/E{total_row}")
    cc.font = FONT_TOTAL; cc.fill = FILL_TOTAL; cc.number_format = NUM_PCT

    # Heatmap on each vendor column
    for j, _ in enumerate(vendors):
        col = 6 + j
        rng = f"{get_column_letter(col)}{start_row}:{get_column_letter(col)}{end_row}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="min", start_color="EFF6FF",
            mid_type="percentile", mid_value=50, mid_color="60A5FA",
            end_type="max", end_color="1E3A8A",
        ))
    pen_col = get_column_letter(7 + len(vendors))
    ws.conditional_formatting.add(f"{pen_col}{start_row}:{pen_col}{end_row}",
        ColorScaleRule(
            start_type="min", start_color="ECFDF5",
            mid_type="percentile", mid_value=50, mid_color="34D399",
            end_type="max", end_color="064E3B",
        ))

    widths = [10, 10, 10, 12, 14] + [16] * len(vendors) + [12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B5"


def build_competitors_json():
    """Pre-bake competitor data for the dashboard."""
    payload = {
        "vendors": {},
        "states": STATES,
    }
    customers = customers_by_vendor()
    for v, meta in COMPETITORS.items():
        cust_list = customers.get(v, [])
        payload["vendors"][v] = {
            **{k: meta[k] for k in ("total", "hq", "hq_state_full", "focus",
                                     "size_target", "notes", "acquired",
                                     "strengths", "weaknesses", "threat")},
            "allocated": vendor_state_total(v),
            "by_state": STATE_PENETRATION[v],
            "named_customers": cust_list,
            "named_customers_count": len(cust_list),
        }
    # Combined per-state
    combined = {st: 0 for st in STATES}
    for v in COMPETITORS:
        for st in STATES:
            combined[st] += STATE_PENETRATION[v][st]
    payload["combined_by_state"] = combined
    payload["addressable_by_state"] = {
        st: STATE_COUNTS[st][0] + STATE_COUNTS[st][1] + STATE_COUNTS[st][2]
        for st in STATES
    }
    payload["state_names"] = STATE_NAMES
    payload["_demo_customers_note"] = (
        "named_customers list is seeded demo data — illustrative only. "
        "Replace by scraping each vendor's case-study page; see "
        "competitor_customers.py for source hints."
    )

    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(payload, f, separators=(",", ":"))
    print(f"Wrote {JSON_PATH}")


def build():
    wb = load_workbook(XLSX_PATH)
    # Remove if already present (idempotent)
    for name in ("Competitor Summary", "Competitor Penetration"):
        if name in wb.sheetnames:
            del wb[name]
    build_summary_tab(wb)
    build_penetration_tab(wb)
    wb.save(XLSX_PATH)
    print(f"Updated {XLSX_PATH}")
    build_competitors_json()


if __name__ == "__main__":
    build()
